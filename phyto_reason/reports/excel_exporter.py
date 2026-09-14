"""
excel_exporter.py — Excel 结果导出器。

使用 openpyxl 生成格式化的 Excel 工作簿。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class ExcelExporter:
    """Excel 表格导出器。"""

    HEADER_FILL = PatternFill(start_color="1f4e79", end_color="1f4e79", fill_type="solid")
    HEADER_FONT = Font(color="ffffff", bold=True, size=11)
    BORDER = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    @staticmethod
    def export_candidates(candidates: list[dict], path: str) -> str:
        """导出候选TF表。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Candidate TFs"

        headers = ["Rank", "Gene ID", "TF Family", "Confidence Level",
                    "Confidence Score", "Evidence Count", "Top Sources"]
        for j, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=j, value=h)
            cell.font = ExcelExporter.HEADER_FONT
            cell.fill = ExcelExporter.HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = ExcelExporter.BORDER

        for i, c in enumerate(candidates[:100], 2):
            top_srcs = list(c.get("top_scores", {}).keys())[:3] if c.get("top_scores") else []
            values = [
                i - 1,
                c.get("gene_id", ""),
                c.get("tf_family", "-"),
                c.get("confidence_level", ""),
                round(c.get("confidence_score", 0), 4),
                c.get("n_evidences", 0),
                ", ".join(top_srcs),
            ]
            for j, v in enumerate(values, 1):
                cell = ws.cell(row=i, column=j, value=v)
                cell.border = ExcelExporter.BORDER

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 14
        ws.column_dimensions["F"].width = 14
        ws.column_dimensions["G"].width = 30

        wb.save(path)
        return path

    @staticmethod
    def export_evidence(evidences: list[dict], path: str) -> str:
        """导出证据详情表。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Evidence Details"

        headers = ["Gene ID", "Source", "Score", "Weight", "Contribution", "Description"]
        for j, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=j, value=h)
            cell.font = ExcelExporter.HEADER_FONT
            cell.fill = ExcelExporter.HEADER_FILL
            cell.border = ExcelExporter.BORDER

        for i, ev in enumerate(evidences[:500], 2):
            values = [
                ev.get("gene_id", ""),
                ev.get("source", ""),
                round(ev.get("normalized_score", 0), 4),
                ev.get("weight", 0),
                round(ev.get("contribution", 0), 4),
                ev.get("explanation", "")[:100],
            ]
            for j, v in enumerate(values, 1):
                cell = ws.cell(row=i, column=j, value=v)
                cell.border = ExcelExporter.BORDER

        for col in ["A", "B", "C", "D", "E"]:
            ws.column_dimensions[col].width = 16
        ws.column_dimensions["F"].width = 50

        wb.save(path)
        return path
